from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from pydantic import BaseModel

from chatppt.app.domain.models.deck_ir import DeckIR
from chatppt.app.domain.models.editable_deck_state import EditableDeckState, EditableSlideState
from chatppt.app.domain.models.render_manifest import RenderManifest, RenderedSlide
from chatppt.app.rendering.layout_matcher import LayoutMatcher
from chatppt.app.rendering.shape_tagger import ShapeTagger
from chatppt.app.rendering.template_loader import TemplateMeta


class RenderOutput(BaseModel):
    pptx_path: Path
    manifest: RenderManifest
    editable_deck_state: EditableDeckState
    deck: DeckIR


class PptxRenderer:
    def __init__(self, layout_matcher: LayoutMatcher | None = None, shape_tagger: ShapeTagger | None = None):
        self.layout_matcher = layout_matcher or LayoutMatcher()
        self.shape_tagger = shape_tagger or ShapeTagger()

    def render(self, deck: DeckIR, template: TemplateMeta, output_root: Path) -> RenderOutput:
        output_root = Path(output_root)
        output_root.mkdir(parents=True, exist_ok=True)
        presentation = Presentation()
        manifest_slides: list[RenderedSlide] = []
        editable_slides: list[EditableSlideState] = []

        for slide in deck.slides:
            matched_layout = self.layout_matcher.match(slide, template)
            layout_meta = next(layout for layout in template.layouts if layout.layout_name == matched_layout.layout_name)
            ppt_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title_box = ppt_slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.0), Inches(0.8))
            title_box.text = slide.title
            self.shape_tagger.tag(title_box, f"{slide.slide_id}:title")

            content_text = "\n".join(block.text for block in slide.content_blocks)
            content_box = ppt_slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.0), Inches(3.2))
            content_box.text = content_text
            self.shape_tagger.tag(content_box, f"{slide.slide_id}:content")

            if slide.visual_slots:
                visual_box = ppt_slide.shapes.add_textbox(Inches(5.2), Inches(1.4), Inches(3.0), Inches(2.5))
                visual_box.text = slide.visual_slots[0].semantic_description
                self.shape_tagger.tag(visual_box, f"{slide.slide_id}:visual")

            placeholder_map = {"title": "title"}
            if slide.content_blocks:
                placeholder_map["content"] = slide.content_blocks[0].block_id
            if slide.visual_slots:
                placeholder_map["visual"] = slide.visual_slots[0].slot_id

            manifest_slides.append(
                RenderedSlide(
                    slide_id=slide.slide_id,
                    layout=matched_layout,
                    placeholder_map=placeholder_map,
                    placeholder_dimensions=layout_meta.placeholder_dimensions,
                    warnings=[],
                )
            )
            editable_slides.append(
                EditableSlideState(
                    slide_id=slide.slide_id,
                    layout_name=matched_layout.layout_name,
                    allowed_edit_operations=slide.allowed_edit_operations,
                    warnings=[],
                )
            )

        pptx_path = output_root / "generated_deck.pptx"
        presentation.save(str(pptx_path))
        return RenderOutput(
            pptx_path=pptx_path,
            manifest=RenderManifest(template_id=template.template_id, slides=manifest_slides),
            editable_deck_state=EditableDeckState(slides=editable_slides),
            deck=deck,
        )
