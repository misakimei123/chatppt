from pathlib import Path

from pptx import Presentation

from chatppt.app.domain.models.deck_ir import DeckIR, SlideIR, SlideContentBlock
from chatppt.app.rendering.pptx_renderer import PptxRenderer
from chatppt.app.rendering.template_loader import TemplateLoader


def test_renderer_outputs_pptx_and_manifest(tmp_path):
    template = TemplateLoader().load(Path("tests/fixtures/templates/default"))
    deck = DeckIR(
        slides=[
            SlideIR(
                slide_id="slide-1",
                section_id="sec-1",
                purpose="title",
                title="Q2 Launch",
                content_blocks=[SlideContentBlock(block_id="b1", kind="bullet", text="Launch goals")],
                visual_slots=[],
                allowed_edit_operations=["rewrite_tone"],
            )
        ]
    )

    result = PptxRenderer().render(deck, template, tmp_path)

    assert result.pptx_path.exists()
    assert result.manifest.slides[0].slide_id == "slide-1"
    presentation = Presentation(str(result.pptx_path))
    assert len(presentation.slides) == 1
